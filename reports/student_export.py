"""Reusable per-student internship report export (CSV / XLSX / PDF).

Builds a report for exactly one student's weekly progress and daily work
logs, scoped to a single internship, and renders it in the requested
format. Authorization is enforced by the calling view; this module only
assembles and renders data.
"""
import csv
import re

from django.http import HttpResponse
from django.utils import timezone


def _fmt_date(value):
    return value.isoformat() if value else "-"


def _fmt_datetime(value):
    if not value:
        return "-"
    return timezone.localtime(value).strftime("%Y-%m-%d %H:%M")


def _reviewed_by_name(obj):
    if obj.reviewed_by_id is None:
        return "-"
    return obj.reviewed_by.get_full_name() or obj.reviewed_by.username


def build_student_report(student, internship):
    """Collect all report data for one student (optionally one internship).

    Returns a dict with student/internship/mentor details plus the weekly
    progress and daily work log rows belonging only to this student.
    """
    from internships.models import FacultyAssignment, SupervisorAssignment
    from progress.models import DailyWorkLog, WeeklyProgress

    weekly_qs = WeeklyProgress.objects.filter(student=student)
    worklog_qs = DailyWorkLog.objects.filter(student=student)
    if internship is not None:
        weekly_qs = weekly_qs.filter(internship=internship)
        worklog_qs = worklog_qs.filter(internship=internship)
    weekly_qs = weekly_qs.select_related("internship", "reviewed_by").order_by("-week_start")
    worklog_qs = worklog_qs.select_related("internship", "reviewed_by").order_by("-date")

    allocation = (
        student.allocations.filter(internship=internship)
        .select_related("internship__company")
        .first()
        if internship is not None
        else None
    )
    mentors = list(
        FacultyAssignment.objects.filter(student=student, active=True)
        .filter(**({"internship": internship} if internship is not None else {}))
        .select_related("faculty", "internship")
    )
    supervisor_assignment = (
        SupervisorAssignment.objects.filter(student=student, active=True)
        .filter(**({"internship": internship} if internship is not None else {}))
        .select_related("supervisor")
        .first()
    )

    weekly_rows = [
        [
            w.week_start.isoformat(),
            w.week_end.isoformat(),
            str(w.progress_percentage) + "%",
            w.get_status_display(),
            w.summary or "-",
            w.completed_tasks or "-",
            w.challenges or "-",
            w.learnings or "-",
            w.next_week_plan or "-",
            _reviewed_by_name(w),
            _fmt_datetime(w.reviewed_at),
            w.review_comments or "-",
        ]
        for w in weekly_qs
    ]
    worklog_rows = [
        [
            log.date.isoformat(),
            str(log.hours_worked),
            log.get_status_display(),
            log.work_description or "-",
            log.tasks_completed or "-",
            log.challenges or "-",
            log.learning_skills or "-",
            _reviewed_by_name(log),
            _fmt_datetime(log.reviewed_at),
            log.review_comments or "-",
        ]
        for log in worklog_qs
    ]

    return {
        "generated_at": timezone.now(),
        "student": student,
        "internship": internship,
        "allocation": allocation,
        "mentors": mentors,
        "supervisor_assignment": supervisor_assignment,
        "weekly_columns": [
            "Week Start", "Week End", "Progress %", "Status", "Summary",
            "Completed Tasks", "Challenges", "Learnings", "Next Week Plan",
            "Reviewed By", "Reviewed At", "Review Comments",
        ],
        "weekly_rows": weekly_rows,
        "worklog_columns": [
            "Date", "Hours Worked", "Status", "Work Description",
            "Tasks Completed", "Challenges", "Learning/Skills",
            "Reviewed By", "Reviewed At", "Review Comments",
        ],
        "worklog_rows": worklog_rows,
    }


def report_has_data(report):
    return bool(report["weekly_rows"] or report["worklog_rows"])


def report_filename(report, ext):
    """StudentName_Internship_Report.<ext>, safe for HTTP headers/filesystems."""
    name = re.sub(r"[^A-Za-z0-9]+", "_", report["student"].full_name).strip("_")
    return f"{name or 'Student'}_Internship_Report.{ext}"


def _summary_pairs(report):
    """Flat label/value pairs describing the student and internship."""
    student = report["student"]
    internship = report["internship"]
    allocation = report["allocation"]
    pairs = [
        ("Student Name", student.full_name),
        ("Student ID / Roll No", student.student_id),
        ("Email", student.email),
        ("Department / Branch", student.department or "-"),
        ("Course", student.course or "-"),
        ("Year", str(student.year) if student.year is not None else "-"),
        ("College", student.college.college_name if student.college else "-"),
    ]
    if internship is not None:
        pairs += [
            ("Internship Title", internship.title),
            ("Company / Organization", internship.company.company_name),
            ("Internship Type", internship.get_internship_type_display()),
            ("Location", internship.location or "-"),
            ("Duration", internship.duration or "-"),
            ("Start Date", _fmt_date(internship.start_date)),
            ("End Date", _fmt_date(internship.end_date)),
            ("Internship Status", internship.get_status_display()),
        ]
    else:
        pairs.append(("Internship", "No active internship"))
    if allocation is not None:
        pairs += [
            ("Allocation Status", allocation.get_status_display()),
            ("Allocation Date", _fmt_date(allocation.allocation_date)),
        ]
    if report["mentors"]:
        pairs.append(
            ("Mentor(s)/Faculty", "; ".join(
                f"{a.faculty.full_name} ({a.faculty.department or '-'})"
                for a in report["mentors"]
            ))
        )
    else:
        pairs.append(("Mentor(s)/Faculty", "-"))
    sa = report["supervisor_assignment"]
    if sa is not None:
        pairs.append(("Company Supervisor", sa.supervisor.name))
    pairs.append(("Report Generated At", _fmt_datetime(report["generated_at"])))
    return pairs


def render_csv(report):
    """Valid CSV with quoted sections; csv.writer escapes , \" and newlines."""
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = (
        f'attachment; filename="{report_filename(report, "csv")}"'
    )
    writer = csv.writer(response)

    def section(title):
        writer.writerow([])
        writer.writerow([title])

    writer.writerow(["Student Internship Report"])
    for label, value in _summary_pairs(report):
        writer.writerow([label, value])

    section("Weekly Progress")
    if report["weekly_rows"]:
        writer.writerow(report["weekly_columns"])
        writer.writerows(report["weekly_rows"])
    else:
        writer.writerow(["No weekly progress records found."])

    section("Daily Worklogs")
    if report["worklog_rows"]:
        writer.writerow(report["worklog_columns"])
        writer.writerows(report["worklog_rows"])
    else:
        writer.writerow(["No daily worklog records found."])
    return response


def render_xlsx(report):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    info = wb.active
    info.title = "Report Info"
    info.append(["Field", "Value"])
    for label, value in _summary_pairs(report):
        info.append([label, value])
    for cell in info[1]:
        cell.font = Font(bold=True)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 60

    def add_table_sheet(title, columns, rows, widths):
        ws = wb.create_sheet(title=title[:31])
        ws.append(columns)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        wrap = Alignment(wrap_text=True, vertical="top")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = wrap
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = width
        ws.freeze_panes = "A2"
        return ws

    if report["weekly_rows"]:
        add_table_sheet(
            "Weekly Progress", report["weekly_columns"], report["weekly_rows"],
            [12, 12, 11, 18, 40, 30, 30, 30, 30, 20, 18, 30],
        )
    else:
        ws = wb.create_sheet(title="Weekly Progress")
        ws.append(["No weekly progress records found."])

    if report["worklog_rows"]:
        add_table_sheet(
            "Daily Worklogs", report["worklog_columns"], report["worklog_rows"],
            [12, 13, 18, 50, 35, 35, 35, 20, 18, 30],
        )
    else:
        ws = wb.create_sheet(title="Daily Worklogs")
        ws.append(["No daily worklog records found."])

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{report_filename(report, "xlsx")}"'
    )
    wb.save(response)
    return response


def render_pdf(report):
    """Clean, paginated A4 report; long text wraps instead of overflowing."""
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="{report_filename(report, "pdf")}"'
    )

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    heading_style = ParagraphStyle(
        "SectionHeading", parent=styles["Heading2"],
        spaceBefore=14, spaceAfter=6, textColor=colors.HexColor("#0d6efd"),
    )
    cell_style = ParagraphStyle(
        "Cell", parent=styles["BodyText"], fontSize=7.5, leading=9.5, spaceAfter=0,
    )
    header_cell_style = ParagraphStyle(
        "HeaderCell", parent=cell_style, textColor=colors.white,
        fontName="Helvetica-Bold",
    )
    muted_style = ParagraphStyle(
        "Muted", parent=styles["BodyText"], textColor=colors.grey,
    )

    def p(text, style):
        # Standard PDF fonts are Latin-1; keep unknown characters readable.
        safe = escape(str(text)).encode("latin-1", "replace").decode("latin-1")
        return Paragraph(safe, style)

    doc = SimpleDocTemplate(
        response, pagesize=landscape(A4),
        leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=15 * mm, bottomMargin=15 * mm,
        title="Student Internship Report",
    )
    elements = [
        p("Student Internship Report", title_style),
        p(
            f"Generated on {_fmt_datetime(report['generated_at'])}", muted_style
        ),
        Spacer(1, 8),
    ]

    summary = Table(
        [[p(k, header_cell_style), p(v, cell_style)] for k, v in _summary_pairs(report)],
        colWidths=[55 * mm, 125 * mm],
    )
    summary.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eef2ff")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(summary)

    def table_section(heading, columns, rows, col_widths):
        elements.append(p(heading, heading_style))
        if not rows:
            elements.append(
                p(f"No {heading.lower()} records found.", muted_style)
            )
            return
        data = [[p(c, header_cell_style) for c in columns]]
        data += [[p(cell, cell_style) for cell in row] for row in rows]
        table = Table(data, colWidths=col_widths, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f7f7f7")]),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        elements.append(table)

    table_section(
        "Weekly Progress", report["weekly_columns"], report["weekly_rows"],
        [17 * mm, 17 * mm, 15 * mm, 20 * mm, 33 * mm, 27 * mm,
         27 * mm, 24 * mm, 24 * mm, 20 * mm, 18 * mm, 25 * mm],
    )
    table_section(
        "Daily Worklogs", report["worklog_columns"], report["worklog_rows"],
        [16 * mm, 14 * mm, 19 * mm, 45 * mm, 32 * mm, 30 * mm,
         30 * mm, 20 * mm, 18 * mm, 23 * mm],
    )

    doc.build(elements)
    return response


RENDERERS = {"csv": render_csv, "xlsx": render_xlsx, "pdf": render_pdf}
