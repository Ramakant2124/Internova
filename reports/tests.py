# reports module

from datetime import timedelta

from django.contrib import messages
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django.utils import timezone

from accounts.factories import (
    create_college,
    create_company,
    create_faculty,
    create_internship,
    create_student,
)
from internships.services import allocate_internship, apply_to_internship, assign_faculty
from reports.views import REPORT_TYPES, get_report_rows


class ReportsTest(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.college = create_college()
        self.company = create_company()
        self.faculty = create_faculty(self.college)
        self.student = create_student(self.college)
        self.internship = create_internship(self.company)
        self.today = timezone.localdate()
        apply_to_internship(self.student, self.internship, "Please.")
        allocate_internship(
            self.student, self.internship, self.college.user,
            self.today - timedelta(days=5), self.today + timedelta(days=20),
        )
        assign_faculty(self.faculty, self.student, self.internship, self.college.user)

    def test_all_report_types_build_rows(self):
        request = self.factory.get("/reports/")
        request.user = self.college.user
        for report_type in REPORT_TYPES:
            columns, rows = get_report_rows(report_type, request)
            self.assertTrue(columns)

    def test_company_scope_filters_rows(self):
        from progress.models import WeeklyProgress

        other_company = create_company(username="other_company")
        other_internship = create_internship(other_company)
        other_student = create_student(self.college, username="other_student")
        WeeklyProgress.objects.create(
            student=self.student, internship=self.internship,
            week_start=self.today - timedelta(days=6), week_end=self.today,
            summary="Mine", progress_percentage=60,
        )
        WeeklyProgress.objects.create(
            student=other_student, internship=other_internship,
            week_start=self.today - timedelta(days=6), week_end=self.today,
            summary="Theirs", progress_percentage=60,
        )
        request = self.factory.get("/reports/")
        request.user = self.company.user
        _, rows = get_report_rows("progress", request)
        self.assertTrue(any(row[0] == self.student.full_name for row in rows))
        self.assertFalse(any(row[0] == other_student.full_name for row in rows))

    def test_report_pages_require_role(self):
        self.client.force_login(self.college.user)
        for report_type in REPORT_TYPES:
            response = self.client.get(reverse("reports:report_detail", args=[report_type]))
            self.assertEqual(response.status_code, 200, report_type)

    def test_export_csv(self):
        self.client.force_login(self.college.user)
        response = self.client.get(
            reverse("reports:report_export", args=["students", "csv"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "text/csv")

    def test_export_xlsx(self):
        self.client.force_login(self.college.user)
        response = self.client.get(
            reverse("reports:report_export", args=["students", "xlsx"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_export_pdf(self):
        self.client.force_login(self.college.user)
        response = self.client.get(
            reverse("reports:report_export", args=["students", "pdf"])
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")

    def test_unknown_report_type_404(self):
        self.client.force_login(self.college.user)
        response = self.client.get(reverse("reports:report_detail", args=["placement"]))
        self.assertEqual(response.status_code, 404)

    def test_students_denied_from_reports(self):
        self.client.force_login(self.student.user)
        response = self.client.get(reverse("reports:home"))
        self.assertEqual(response.status_code, 403)

    def test_weekly_progress_populates_progress_report(self):
        from progress.models import WeeklyProgress

        WeeklyProgress.objects.create(
            student=self.student, internship=self.internship,
            week_start=self.today - timedelta(days=6), week_end=self.today,
            summary="Week of work.", progress_percentage=50,
        )
        request = self.factory.get("/reports/")
        request.user = self.college.user
        _, rows = get_report_rows("progress", request)
        self.assertTrue(any(self.student.full_name in row for row in rows))


class StudentReportExportTest(TestCase):
    """Per-student report export from the dashboards (CSV/XLSX/PDF)."""

    def setUp(self):
        self.college = create_college()
        self.company = create_company()
        self.faculty = create_faculty(self.college)
        self.student = create_student(self.college)
        self.internship = create_internship(self.company)
        self.today = timezone.localdate()
        apply_to_internship(self.student, self.internship, "Please.")
        allocate_internship(
            self.student, self.internship, self.college.user,
            self.today - timedelta(days=5), self.today + timedelta(days=20),
        )
        assign_faculty(self.faculty, self.student, self.internship, self.college.user)

    @staticmethod
    def _url(student, fmt):
        return reverse("reports:student_report_export", args=[student.pk, fmt])

    @staticmethod
    def _weekly(student, internship, **extra):
        from progress.models import WeeklyProgress

        fields = dict(
            student=student, internship=internship,
            week_start=timezone.localdate() - timedelta(days=6),
            week_end=timezone.localdate(),
            summary="Built the login flow.",
            completed_tasks="Auth module",
            challenges="Token refresh",
            learnings="JWT basics",
            next_week_plan="Payments",
            progress_percentage=45,
        )
        fields.update(extra)
        return WeeklyProgress.objects.create(**fields)

    @staticmethod
    def _worklog(student, internship, **extra):
        from progress.models import DailyWorkLog

        return DailyWorkLog.objects.create(
            student=student, internship=internship,
            date=timezone.localdate(), work_description="Fixed bugs, wrote tests",
            hours_worked=6, tasks_completed="Issue #12",
            challenges="Flaky test", learning_skills="pytest fixtures", **extra,
        )

    def test_college_can_export_all_formats(self):
        self._weekly(self.student, self.internship)
        self._worklog(self.student, self.internship)
        self.client.force_login(self.college.user)
        expected = {
            "csv": ("text/csv", f'"{self.student.full_name}_Internship_Report.csv"'),
            "xlsx": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                f'"{self.student.full_name}_Internship_Report.xlsx"',
            ),
            "pdf": ("application/pdf", f'"{self.student.full_name}_Internship_Report.pdf"'),
        }
        for fmt, (content_type, disposition) in expected.items():
            response = self.client.get(self._url(self.student, fmt))
            self.assertEqual(response.status_code, 200, fmt)
            self.assertTrue(response["Content-Type"].startswith(content_type), fmt)
            self.assertEqual(response["Content-Disposition"], f"attachment; filename={disposition}")

    def test_faculty_can_export_assigned_student(self):
        self._weekly(self.student, self.internship)
        self.client.force_login(self.faculty.user)
        response = self.client.get(self._url(self.student, "csv"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Weekly Progress", content)
        self.assertIn("Built the login flow.", content)

    def test_company_can_export_allocated_student(self):
        self._worklog(self.student, self.internship)
        self.client.force_login(self.company.user)
        response = self.client.get(self._url(self.student, "csv"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Daily Worklogs", content)
        self.assertIn("Fixed bugs, wrote tests", content)

    def test_csv_contains_only_selected_student_data(self):
        other_student = create_student(self.college, username="other_student")
        self._weekly(self.student, self.internship, summary="Alice data")
        self._weekly(other_student, self.internship, summary="Bob data")
        self.client.force_login(self.college.user)
        response = self.client.get(self._url(self.student, "csv"))
        content = response.content.decode()
        self.assertIn("Alice data", content)
        self.assertNotIn("Bob data", content)

    def test_pdf_and_xlsx_contain_only_selected_student_data(self):
        other_student = create_student(self.college, username="other_student")
        self._weekly(self.student, self.internship, summary="Mine only")
        self._weekly(other_student, self.internship, summary="Theirs only")
        self.client.force_login(self.college.user)
        pdf_response = self.client.get(self._url(self.student, "pdf"))
        self.assertEqual(pdf_response.status_code, 200)
        self.assertGreater(len(pdf_response.content), 1000)
        xlsx_response = self.client.get(self._url(self.student, "xlsx"))
        self.assertEqual(xlsx_response.status_code, 200)
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(xlsx_response.content))
        texts = [
            str(cell.value)
            for sheet in wb.worksheets
            for row in sheet.iter_rows()
            for cell in row
        ]
        joined = "\n".join(texts)
        self.assertIn(self.student.full_name, joined)
        self.assertNotIn(other_student.full_name, joined)
        self.assertIn("Mine only", joined)
        self.assertNotIn("Theirs only", joined)

    def test_faculty_cannot_export_unassigned_student(self):
        outsider = create_student(self.college, username="outsider")
        self._weekly(outsider, self.internship)
        self.client.force_login(self.faculty.user)
        response = self.client.get(self._url(outsider, "csv"))
        self.assertEqual(response.status_code, 403)

    def test_company_cannot_export_other_company_student(self):
        other_company = create_company(username="other_company")
        other_internship = create_internship(other_company)
        other_student = create_student(self.college, username="other_co_student")
        apply_to_internship(other_student, other_internship, "Please.")
        allocate_internship(
            other_student, other_internship, self.college.user,
            self.today - timedelta(days=5), self.today + timedelta(days=20),
        )
        self._weekly(other_student, other_internship)
        self.client.force_login(self.company.user)
        response = self.client.get(self._url(other_student, "csv"))
        self.assertEqual(response.status_code, 403)

    def test_student_role_denied(self):
        self._weekly(self.student, self.internship)
        self.client.force_login(self.student.user)
        response = self.client.get(self._url(self.student, "csv"))
        self.assertEqual(response.status_code, 403)

    def test_unknown_format_404(self):
        self._weekly(self.student, self.internship)
        self.client.force_login(self.college.user)
        response = self.client.get(self._url(self.student, "docx"))
        self.assertEqual(response.status_code, 404)

    def test_no_records_redirects_with_message_for_every_role(self):
        cases = (
            ("college", self.college.user),
            ("faculty", self.faculty.user),
            ("company", self.company.user),
        )
        for template_dir, user in cases:
            with self.subTest(role=template_dir):
                self.client.force_login(user)
                response = self.client.get(self._url(self.student, "pdf"), follow=True)
                self.assertEqual(response.status_code, 200)
                self.assertTemplateUsed(response, f"{template_dir}/student_detail.html")
                messages_list = list(messages.get_messages(response.wsgi_request))
                self.assertTrue(any("No weekly progress" in str(m) for m in messages_list))

    def test_only_worklog_still_exports(self):
        self._worklog(self.student, self.internship)
        self.client.force_login(self.college.user)
        response = self.client.get(self._url(self.student, "csv"))
        self.assertEqual(response.status_code, 200)
        content = response.content.decode()
        self.assertIn("Daily Worklogs", content)
        self.assertIn("No weekly progress records found.", content)

    def test_special_characters_survive_csv(self):
        self._weekly(self.student, self.internship, summary='Commas, "quotes"\nand line breaks')
        self.client.force_login(self.college.user)
        response = self.client.get(self._url(self.student, "csv"))
        content = response.content.decode()
        self.assertIn('"Commas, ""quotes""', content)
        self.assertIn("and line breaks", content)
